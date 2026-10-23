from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os
D=os.path.dirname(os.path.abspath(__file__)); I=lambda n:os.path.join(D,'img',n)
GREEN=RGBColor(0x14,0x68,0x4A); INK=RGBColor(0x0D,0x1F,0x17); MUT=RGBColor(0x4A,0x5A,0x52); WHITE=RGBColor(255,255,255)
doc=Document(); sec=doc.sections[0]
sec.page_width=Mm(210); sec.page_height=Mm(297)
for a in('left_margin','right_margin'): setattr(sec,a,Mm(15))
sec.top_margin=Mm(12); sec.bottom_margin=Mm(16)
st=doc.styles['Normal']; st.font.name='Arial'; st.font.size=Pt(10); st.element.rPr.rFonts.set(qn('w:eastAsia'),'Arial')
st.paragraph_format.space_after=Pt(0)
W=180 # mm usable

def shade(cell,hex):
    tcPr=cell._tc.get_or_add_tcPr(); s=OxmlElement('w:shd'); s.set(qn('w:val'),'clear'); s.set(qn('w:color'),'auto'); s.set(qn('w:fill'),hex); tcPr.append(s)
def margins(cell,t=0,l=0,b=0,r=0):
    tcPr=cell._tc.get_or_add_tcPr(); m=OxmlElement('w:tcMar')
    for k,v in(('top',t),('left',l),('bottom',b),('right',r)):
        e=OxmlElement('w:'+k); e.set(qn('w:w'),str(int(v*56.7))); e.set(qn('w:type'),'dxa'); m.append(e)
    tcPr.append(m)
def noborders(tbl):
    tblPr=tbl._tbl.tblPr; b=OxmlElement('w:tblBorders')
    for k in('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+k); e.set(qn('w:val'),'nil'); b.append(e)
    tblPr.append(b)
def table(rows,cols,widths):
    t=doc.add_table(rows=rows,cols=cols); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False; noborders(t)
    lay=OxmlElement('w:tblLayout'); lay.set(qn('w:type'),'fixed'); t._tbl.tblPr.append(lay)
    for col,w in zip(t.columns,widths): col.width=Mm(w)
    for r in t.rows:
        for c,w in zip(r.cells,widths): c.width=Mm(w)
    return t
def para(container,text='',size=10,bold=False,color=INK,before=0,after=0,spacing=None,align=None,line=None):
    p=container.add_paragraph() if hasattr(container,'add_paragraph') else container
    p.paragraph_format.space_before=Pt(before); p.paragraph_format.space_after=Pt(after)
    if line: p.paragraph_format.line_spacing=line
    if align: p.alignment=align
    if text:
        r=p.add_run(text); r.font.size=Pt(size); r.bold=bold; r.font.color.rgb=color
        if spacing: r._r.get_or_add_rPr().append(_sp(spacing))
    return p
def _sp(v):
    e=OxmlElement('w:spacing'); e.set(qn('w:val'),str(v)); return e
def first(cell,text='',**k):
    p=cell.paragraphs[0]; return para(p,text,**k)
def eyebrow(c,text,color=GREEN,first_=False,before=0):
    return (first(c,text,size=8,bold=True,color=color,spacing=30,before=before) if first_ else para(c,text,size=8,bold=True,color=color,spacing=30,before=before))
def rule(c,before=4,after=8):
    p=para(c,'',before=before,after=after)
    pPr=p._p.get_or_add_pPr(); b=OxmlElement('w:pBdr'); e=OxmlElement('w:top')
    e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'24'); e.set(qn('w:color'),'2EE07F'); e.set(qn('w:space'),'1'); b.append(e); pPr.append(b)
    p.paragraph_format.right_indent=Mm(150)
def pic(c_or_p,name,w,align=None):
    p=c_or_p.add_paragraph() if hasattr(c_or_p,'add_paragraph') else c_or_p
    if align: p.alignment=align
    p.add_run().add_picture(I(name),width=Mm(w)); return p
def pagebreak(): doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
def footer(n):
    p=para(doc,'',before=0); 
    p.add_run(f'DYN HUB, S.L. · www.dynhub.es\t\t{n} / 4').font.size=Pt(7.5)
    for r in p.runs: r.font.color.rgb=MUT
def heading(c,eyebrow_t,title,first_=True):
    eyebrow(c,eyebrow_t,first_=first_)
    p=para(c,title,size=24,bold=True,before=4,line=1.0); rule(c)
def grid(items,cols,icon_less=True,title_size=10):
    t=table(-(-len(items)//cols),cols,[W/cols]*cols)
    for i,(h,body) in enumerate(items):
        c=t.cell(i//cols,i%cols); margins(c,2,0,6,4)
        first(c,h,size=title_size,bold=True); para(c,body,size=9,color=MUT,before=2,line=1.25)
    return t

# ---------- PAGE 1
t=table(1,1,[W]); c=t.cell(0,0); shade(c,'1F9D63'); margins(c,8,8,8,8)
r=table.__class__  # noqa
first(c,'',after=0).add_run().add_picture(I('logo_cat_t.png'),height=Mm(9))
eyebrow(c,'GESTIÓN DE COLONIAS FELINAS URBANAS',WHITE,before=14)
para(c,'Un programa municipal que se puede demostrar.',size=30,bold=True,color=WHITE,before=8,line=1.0)
para(c,'Censo, fichas clínicas, control CER e incidencias en una sola plataforma',size=11,color=WHITE,before=10)
para(c,'LISTO PARA HOY.',size=12,bold=True,color=WHITE,before=6)
pic(c,'shot1.png',100,WD_ALIGN_PARAGRAPH.RIGHT)
pic(c,'logo_dyn_t.png',34).paragraph_format.space_before=Pt(4)
para(doc,'',size=4)
t=table(1,4,[W/4]*4)
for i,(n,l) in enumerate([('6','CAPAS DE MAPA'),('4','IDIOMAS NATIVOS'),('100%','WEB + MÓVIL'),('24/7','ACCESO EN CAMPO')]):
    c=t.cell(0,i); shade(c,'F1F1F1'); margins(c,5,6,5,4)
    first(c,n,size=15,bold=True,color=GREEN); para(c,l,size=7,bold=True,color=MUT,spacing=10)
c=doc.add_table(rows=1,cols=1); noborders(c); c=c.cell(0,0); margins(c,8,0,0,0)
eyebrow(c,'EL PROBLEMA',first_=True)
para(c,'La mayoría de ayuntamientos gestionan el programa con hojas de cálculo y grupos de mensajería. Cuando llega una auditoría o una pregunta política, nadie puede responder en minutos.',size=10,before=6)
rule(c,6,6)
p=para(c,'',size=10); r=p.add_run('CATMetrix'); r.bold=True; p.add_run(' convierte esta obligación legal en un procedimiento digital, trazable y acreditable desde el primer día.')
para(c,'',before=4)
grid([('Alta en días','Configuración del municipio y catálogos iniciales, sin desarrollo a medida.'),
      ('Migración incluida','Su hoja de cálculo actual, cargada y validada antes de escribir nada en el censo.'),
      ('Formación por perfiles','Sesiones diferenciadas para personal técnico, asociaciones y voluntariado.')],3)
footer(1); pagebreak()

# ---------- PAGE 2
t=table(1,1,[W]); heading(t.cell(0,0),'FUNCIONALIDADES','Seis módulos, un solo registro.')
grid([('Censo territorial','Cartografía de seis capas sobre el IGN: colonias, incidencias, zonas sensibles y radios de campeo.'),
      ('Ficha clínica','Historial completo por gato, con certificados oficiales de esterilización y vacunación en PDF.'),
      ('Sistema CER','Captura, esterilización y retorno con indicadores de eficacia y alertas de ciclos estancados.'),
      ('Incidencias vecinales','Quejas geolocalizadas con flujo de estados y expediente acreditativo completo.'),
      ('Voluntariado y asociaciones','Credencial digital, validación municipal de propuestas y mensajería interna por ámbitos.'),
      ('Cuadro de mando','Informes automáticos, periódicos y exportables, con las mismas cifras que la pantalla.')],2)
para(doc,'',size=6)
t=table(2,3,[W/3]*3)
for i,(im,h,b) in enumerate([('map.png','Mapa municipal','Colonias, incidencias y zonas sensibles sobre cartografía del IGN.'),
    ('dash.png','Cuadro de mando','Indicadores generales y evolución mensual del programa.'),
    ('cert.png','Certificado oficial','Generado en PDF con el escudo y el sello del Ayuntamiento.')]):
    c=t.cell(0,i); margins(c,0,0,0,5); first(c).add_run().add_picture(I(im),width=Mm(55))
    c=t.cell(1,i); margins(c,2,0,0,5); first(c,h,size=8.5,bold=True,color=GREEN); para(c,b,size=8,color=MUT)
para(doc,'',size=8)
t=table(1,2,[85,95]); c=t.cell(0,0); margins(c,0,0,0,6); first(c).add_run().add_picture(I('export.png'),width=Mm(80))
c=t.cell(0,1); margins(c,12,6,0,0)
first(c,'Exportación e interoperabilidad',size=13,bold=True,color=GREEN)
para(c,'Cualquier entidad, o el volcado completo del programa, se exporta a Excel y CSV. El censo se publica también en GeoJSON estándar a través de la API, listo para cualquier sistema de información geográfica municipal.',size=9.5,before=6,line=1.3)
para(doc,'',size=30); footer(2); pagebreak()

# ---------- PAGE 3
t=table(1,1,[W]); c=t.cell(0,0); heading(c,'TRABAJO EN CAMPO','La misma aplicación, lista para la calle.')
para(c,'La aplicación móvil, disponible para Android y para iOS, comparte con la web el mismo modelo de datos y las mismas reglas. No es una versión reducida: es la misma herramienta, pensada para el trabajo en el punto exacto donde ocurre.',size=10,color=MUT,after=10,line=1.3)
grid([('Geolocalización real','La ubicación la toma el mismo dispositivo al dar de alta una colonia, un punto de alimentación o un aviso.'),
      ('Fotografía desde la cámara','Para la ficha del animal y para el recinto de la colonia, sin pasos intermedios.'),
      ('Mismo registro que la web','Colonias, gatos, CER, incidencias y voluntariado, dados de alta directamente en el punto de trabajo.')],3)
para(doc,'',size=8)
t=table(2,3,[W/3]*3)
for i,(im,h) in enumerate([('ph1.png','Mapa con capas activables'),('ph2.png','Ficha clínica con cámara'),('ph3.png','Sistema CER')]):
    c=t.cell(0,i); margins(c,0,3,0,3); pic(c.paragraphs[0],im,52,WD_ALIGN_PARAGRAPH.CENTER)
    c=t.cell(1,i); first(c,h,size=8,bold=True,color=RGBColor(0x44,0x44,0x44),before=3,align=WD_ALIGN_PARAGRAPH.CENTER)
c=doc.add_table(rows=1,cols=1); noborders(c); c=c.cell(0,0); margins(c,6,0,0,40); rule(c,6,6)
first(c,'El efecto sobre la calidad del censo es inmediato: el dato se registra donde y cuando se observa, que es el único momento en que es exacto.',size=10,bold=True)
para(doc,'',size=40); footer(3); pagebreak()

# ---------- PAGE 4
t=table(1,1,[W]); heading(t.cell(0,0),'POR QUÉ DYN HUB','Producto propio. Sin intermediarios.')
t=table(4,2,[14,166])
for i,(h,b) in enumerate([('Arquitectura multimunicipio','Cada ayuntamiento configura su propio entorno: identidad, catálogos y encuadre, sin tocar código.'),
    ('Multiidioma nativo','Catalán, castellano, euskera y gallego, también en los certificados e informes en PDF.'),
    ('Seguridad y RGPD','Autenticación obligatoria, cifrado en tránsito, copias diarias y trazabilidad de autoría en cada acción.'),
    ('Evolución continua','Mantenimiento correctivo, adaptativo y evolutivo incluido durante toda la vigencia del contrato.')]):
    c=t.cell(i,0); margins(c,2,0,6,0); first(c,f'0{i+1}',size=12,bold=True,color=RGBColor(0x8A,0xD7,0xB0))
    c=t.cell(i,1); margins(c,2,0,6,0); first(c,h,size=10,bold=True); para(c,b,size=9,color=MUT,before=2)
c=doc.add_table(rows=1,cols=1); noborders(c); c=c.cell(0,0); margins(c,6,0,0,0)
eyebrow(c,'CUMPLIMIENTO NORMATIVO',first_=True)
t=table(1,3,[45,45,70])
for i,x in enumerate(['Ley 7/2023 · art. 39','RGPD · LO 3/2018','RD 1112/2018 · WCAG 2.1 AA']):
    c=t.cell(0,i); shade(c,'E9F7EF'); margins(c,3,4,3,4); first(c,x,size=8,bold=True,color=GREEN)
para(doc,'',size=6)
c=doc.add_table(rows=1,cols=1); noborders(c); c=c.cell(0,0); eyebrow(c,'Y ADEMÁS',first_=True)
t=table(1,2,[88,92])
for i,(h,b) in enumerate([('Catalán · Castellano · Euskera · Gallego','También en certificados e informes.'),('Migración incluida','Su hoja de cálculo actual, cargada y validada.')]):
    c=t.cell(0,i); shade(c,'F1F1F1'); margins(c,5,5,5,5); first(c,h,size=9.5,bold=True); para(c,b,size=9,color=MUT,before=2)
para(doc,'',size=110)
t=table(1,2,[110,70]); c=t.cell(0,0); shade(c,'1F9D63'); margins(c,8,8,8,8)
first(c,'Hablemos.',size=16,bold=True,color=WHITE)
para(c,'Le ofrecemos una demostración sin compromiso sobre datos reales de su municipio.',size=9,color=WHITE,before=4)
para(c,'WEB: www.dynhub.es   ·   TELÉFONO: 620 082 395   ·   EMAIL: ignacio.moure@dynhub.es',size=8.5,bold=True,color=WHITE,before=8)
c=t.cell(0,1); shade(c,'1F9D63'); margins(c,10,8,8,8); pic(c.paragraphs[0],'logo_dyn2_t.png',42,WD_ALIGN_PARAGRAPH.RIGHT)
doc.save(os.path.join(D,'Ficha_Producto_CATMetrix_ES.docx'))
